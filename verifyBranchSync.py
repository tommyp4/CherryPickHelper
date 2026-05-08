import os
from git import Repo
import config
from openpyxl import load_workbook

def main():
    if not os.path.exists(config.local_repo_path):
        print(f"❌ Error: Repository path {config.local_repo_path} not found.")
        return

    excel_file = 'cherrypick_list.xlsx'
    if not os.path.exists(excel_file):
        print(f"❌ Error: {excel_file} not found. Run pullRequestHelper.py first.")
        return

    repo = Repo(config.local_repo_path)
    git = repo.git

    print(f"--- Definitive Branch Sync Verification ---")
    print(f"Base (Release):   {config.cherry_pick_base_branch}")
    print(f"Source (Develop): {config.main_target_branch}")
    print("-" * 40)

    print("Fetching latest changes from remote...")
    try:
        git.execute(['git', 'remote', 'prune', 'origin'])
    except: pass
    repo.remotes.origin.fetch()

    # Use origin versions to be 100% sure we are comparing against the server state
    upstream = f"origin/{config.cherry_pick_base_branch}"
    head = f"origin/{config.main_target_branch}"

    print(f"Analyzing patches between {upstream} and {head}...")
    try:
        # git cherry -v <upstream> <head>
        # + means the patch is missing from upstream
        # - means the patch is already in upstream
        cherry_output = git.cherry("-v", upstream, head)
    except Exception as e:
        print(f"❌ Error running git cherry: {e}")
        return

    # Map commit IDs to their sync status according to git cherry
    # status_map[full_hash] = 'Yes' (if -) or 'No' (if +)
    status_map = {}
    lines = cherry_output.splitlines()
    for line in lines:
        parts = line.split()
        if len(parts) >= 2:
            status = 'Yes' if parts[0] == '-' else 'No'
            commit_id = parts[1]
            status_map[commit_id] = status

    print(f"Scanned {len(status_map)} commits in history.")
    print(f"Updating {excel_file} with patch verification...")

    try:
        wb = load_workbook(excel_file)
        ws = wb.active

        # 1. Find columns
        col_map = {cell.value: cell.column for cell in ws[1]}
        if 'Commit ID' not in col_map:
            print("❌ Error: 'Commit ID' column not found in Excel.")
            return

        # Ensure 'Patch physically in Release?' column exists
        col_name = 'Patch physically in Release?'
        if col_name not in col_map:
            new_col = ws.max_column + 1
            ws.cell(row=1, column=new_col).value = col_name
            target_col = new_col
        else:
            target_col = col_map[col_name]

        id_col = col_map['Commit ID']

        # 2. Iterate rows and match by Commit ID
        update_count = 0
        for row_idx in range(2, ws.max_row + 1):
            commit_id = ws.cell(row=row_idx, column=id_col).value
            
            # Check for match (handling potential short vs full hashes)
            # git cherry usually gives full hashes
            found_status = None
            if commit_id in status_map:
                found_status = status_map[commit_id]
            else:
                # Fallback: check if any key in status_map starts with our commit_id
                for full_hash, status in status_map.items():
                    if full_hash.startswith(str(commit_id)) or str(commit_id).startswith(full_hash):
                        found_status = status
                        break
            
            if found_status:
                ws.cell(row=row_idx, column=target_col).value = found_status
                update_count += 1

        wb.save(excel_file)
        print(f"✅ Successfully updated {update_count} rows in {excel_file}.")

    except PermissionError:
        print(f"❌ Error: Could not save to {excel_file}. Please close the file.")
    except Exception as e:
        print(f"❌ Error updating Excel: {e}")

if __name__ == "__main__":
    main()
