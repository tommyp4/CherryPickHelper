import os, re
from jira import JIRA
from openpyxl import load_workbook
import config

def get_jira_decision(versions, current_status, hotfix_set):
    """Returns (decision, updated_status) based on Jira Fix Versions and git state."""
    is_physically_present = current_status.startswith("Yes")
    is_ambiguous = current_status.startswith("Likely") or current_status.startswith("Needs attention")

    if is_physically_present:
        return 'no', None
    if is_ambiguous:
        return '', None
    if not versions:
        return 'no', None
    if any(v == config.jira_branched_from_version for v in versions):
        if current_status == "No":
            return 'yes', "No (⚠️ Missed in previous release?)"
        return 'no', None
    if any(v == config.jira_target_version for v in versions):
        return 'yes', None
    if any(v in hotfix_set for v in versions):
        return 'yes', None
    return 'no', None

def extract_jira_id_from_link(link_formula):
    """Extracts the ID from an Excel HYPERLINK formula."""
    if not link_formula or not isinstance(link_formula, str):
        return None
    # Matches HYPERLINK("...","ID") or just ID if it's not a formula
    match = re.search(r'ALP[-_]?(\d+)', link_formula, re.IGNORECASE)
    if match:
        return f"ALP-{match.group(1)}"
    return None

def main():
    # === JIRA AUTHENTICATION ===
    jira_server = config.jira_base_url.split('/browse/')[0]
    jira_user = config.jira_email
    jira_token = config.jira_api_token

    if not jira_user or jira_user == 'your_email@company.com':
        print("❌ Error: jira_email not set correctly in config.py.")
        return

    print(f"Connecting to Jira at {jira_server}...")
    try:
        jira = JIRA(server=jira_server, basic_auth=(jira_user, jira_token))
    except Exception as e:
        print(f"❌ Connection error: {e}")
        return

    # === CONFIGURATION ===
    excel_file = 'cherrypick_list.xlsx'
    if not os.path.exists(excel_file):
        print(f"{excel_file} not found. Run pullRequestHelper.py first.")
        return

    print(f"Reading {excel_file} to look up Fix Versions...")
    
    try:
        wb = load_workbook(excel_file)
        ws = wb.active

        # 1. Find relevant columns
        col_map = {cell.value: cell.column for cell in ws[1]}
        required = ['Jira ID', 'Audit Trail', 'Action']
        if not all(col in col_map for col in required):
            print(f"❌ Error: Required columns {required} not found.")
            return

        # Ensure 'Jira Target' column exists
        if 'Jira Target' not in col_map:
            new_col = ws.max_column + 1
            ws.cell(row=1, column=new_col).value = 'Jira Target'
            fix_version_col = new_col
        else:
            fix_version_col = col_map['Jira Target']

        jira_col = col_map['Jira ID']
        status_col = col_map['Audit Trail']
        decision_col = col_map['Action']
        
        # 2. Cache lookups
        cache = {} 

        # 3. Iterate rows (starting at row 2)
        total_rows = ws.max_row - 1
        print(f"Checking {total_rows} rows...")

        hotfix_set = set(config.jira_hotfix_versions)

        for row_idx in range(2, ws.max_row + 1):
            cell_value = ws.cell(row=row_idx, column=jira_col).value
            jira_id = extract_jira_id_from_link(cell_value)
            
            current_status = str(ws.cell(row=row_idx, column=status_col).value)

            version_str = "No Fix Version"
            if jira_id:
                if jira_id in cache:
                    version_str = cache[jira_id]
                    versions = [v.strip() for v in version_str.split(',')] if version_str != "No Fix Version" else []
                    decision, new_status = get_jira_decision(versions, current_status, hotfix_set)
                    ws.cell(row=row_idx, column=decision_col).value = decision
                    if new_status:
                        ws.cell(row=row_idx, column=status_col).value = new_status
                else:
                    try:
                        print(f"  [{row_idx-1}/{total_rows}] Fetching {jira_id}...", end="\r")
                        issue = jira.issue(jira_id, fields='fixVersions')
                        versions = [v.name for v in issue.fields.fixVersions]
                        
                        if versions:
                            version_str = ", ".join(versions)

                        decision, new_status = get_jira_decision(versions, current_status, hotfix_set)
                        ws.cell(row=row_idx, column=decision_col).value = decision
                        if new_status:
                            ws.cell(row=row_idx, column=status_col).value = new_status
                        
                        cache[jira_id] = version_str
                    except Exception:
                        version_str = "Error/Not Found"
                
                ws.cell(row=row_idx, column=fix_version_col).value = version_str
            else:
                ws.cell(row=row_idx, column=fix_version_col).value = "N/A"

        wb.save(excel_file)
        print(f"\n✅ Successfully updated {excel_file} with Jira Fix Versions and final decisions.")

    except Exception as e:
        print(f"\n❌ Error processing Excel: {e}")

if __name__ == "__main__":
    main()
