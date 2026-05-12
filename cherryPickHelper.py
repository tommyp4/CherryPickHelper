import os, base64, argparse
import pandas as pd
from git import Repo, Git, GitCommandError
import config
from openpyxl import load_workbook

def main():
    # === ARGUMENT PARSING ===
    parser = argparse.ArgumentParser(description="Automated Cherry-Pick Helper")
    parser.add_argument("--blanks", choices=['y', 'n'], help="Cherry-pick ALL blanks? (y/n)")
    parser.add_argument("--pause", choices=['y', 'n'], help="Pause at conflicts? (y/n)")
    args = parser.parse_args()

    # === CONFIGURATION ===
    excel_file = 'cherrypick_list.xlsx'
    
    if not os.path.exists(excel_file):
        print(f"{excel_file} not found. Run pullRequestHelper.py first.")
        return

    print(f"Reading commits from {excel_file}...")
    df = pd.read_excel(excel_file)

    if 'Action' not in df.columns:
        print(f"Error: Column 'Action' not found in {excel_file}")
        return

    # === GLOBAL PROMPTS / ARGUMENTS ===
    # 1. Blanks handling
    has_blanks = df['Action'].isna().any() or (df['Action'].astype(str).str.lower().str.strip() == '').any()
    cherry_pick_blanks = False
    
    if has_blanks:
        if args.blanks:
            user_choice = args.blanks
        else:
            user_choice = input("\n❓ Found blank entries in 'Action' column. Cherry-pick ALL blanks? (y/n): ").lower().strip()
        
        if user_choice == 'y':
            print("✅ Will cherry-pick blanks.")
            cherry_pick_blanks = True
        else:
            print("⏭️  Will skip blanks.")

    # 2. Conflict handling
    if args.pause:
        pause_choice = args.pause
    else:
        pause_choice = input("\n❓ Pause at conflicts for manual resolution? (y/n): ").lower().strip()
    
    pause_at_conflicts = (pause_choice == 'y')
    if pause_at_conflicts:
        print("⏸️  Will pause at conflicts.")
    else:
        print("🤖 Will automatically skip conflicts (Unattended Mode).")

    # Identify commits to cherry-pick
    to_pick_indices = []
    for idx, row in df.iterrows():
        val = str(row['Action']).lower().strip()
        is_blank = pd.isna(row['Action']) or val == '' or val == 'nan'
        if val == 'yes' or (is_blank and cherry_pick_blanks):
            to_pick_indices.append(idx)

    if not to_pick_indices:
        print("No commits marked for cherry-picking.")
        return

    print(f"Found {len(to_pick_indices)} commits to process.")

    # === PREPARE REPOSITORY ===
    repo_url = f'{config.organization_url}/{config.project_name}/_git/{config.repository_name}'
    _pat_b64 = base64.b64encode(f':{config.personal_access_token}'.encode()).decode()
    _auth_header = f'http.extraheader=AUTHORIZATION: Basic {_pat_b64}'

    if not os.path.exists(config.local_repo_path):
        print(f"Cloning repository to {config.local_repo_path}...")
        Git().execute(['git', '-c', _auth_header, 'clone', repo_url, config.local_repo_path])
    else:
        print(f"Using existing repository at {config.local_repo_path}")

    repo = Repo(config.local_repo_path)
    if config.personal_access_token in repo.remotes.origin.url:
        repo.remotes.origin.set_url(repo_url)
    git = repo.git

    # === ROBUST CLEANUP ===
    print("Performing pre-flight cleanup...")
    try:
        git.execute(['git', 'remote', 'prune', 'origin'])
    except GitCommandError: pass
    try:
        git.execute(['git', 'cherry-pick', '--abort'])
    except GitCommandError: pass
    try:
        git.execute(['git', 'merge', '--abort'])
    except GitCommandError: pass
    
    print("Clearing local changes and index errors...")
    git.reset('--hard')
    git.clean('-fd')

    print("Fetching all changes from remote...")
    git.execute(['git', '-c', _auth_header, 'fetch', 'origin'])

    base_branch = config.cherry_pick_base_branch
    target_branch = config.target_branch
    
    print(f"Updating base branch '{base_branch}' from origin...")
    git.checkout('-f', base_branch)
    git.reset('--hard', f'origin/{base_branch}')

    print(f"Ensuring local branch '{target_branch}' is clean...")
    try:
        if repo.active_branch.name == target_branch: git.checkout(base_branch)
        if target_branch in repo.heads: git.branch('-D', target_branch)
    except Exception:
        pass

    print(f"Creating new branch '{target_branch}' from '{base_branch}'...")
    git.checkout('-b', target_branch, base_branch)

    # === CHERRY-PICK LOOP ===
    results = {} # { commit_id: status_string }
    success_count = 0
    fail_count = 0
    empty_count = 0

    for i, idx in enumerate(to_pick_indices):
        row = df.loc[idx]
        commit_id = row['Commit SHA']
        msg = row['Description']
        
        print(f"\n[{i + 1}/{len(to_pick_indices)}] Cherry-picking {commit_id[:8]} - {str(msg)[:50]}...")
        
        try:
            git.cherry_pick(commit_id)
            print("✅ Success")
            results[commit_id] = 'yes'
            success_count += 1
        except GitCommandError as e:
            error_output = str(e).lower()
            if "empty" in error_output or "nothing to commit" in error_output:
                print("ℹ️  No changes to commit (already present).")
                results[commit_id] = 'no changes to commit'
                empty_count += 1
                try: git.execute(['git', 'cherry-pick', '--abort'])
                except GitCommandError: pass
            else:
                if pause_at_conflicts:
                    print(f"\n❌ Conflict in {commit_id[:8]}.")
                    print("🛠️  Resolve in VS/Git and Stage changes.")
                    choice = input("👉 Enter 'r' to continue, or 's' to skip/abort this commit: ").lower().strip()
                    if choice == 'r':
                        try:
                            git.execute(['git', 'cherry-pick', '--continue'], env={'GIT_EDITOR': 'true'})
                            print("✅ Success (Resolved Manually)")
                            results[commit_id] = 'yes'
                            success_count += 1
                        except GitCommandError:
                            print(f"❌ Resolution failed. Skipping.")
                            results[commit_id] = 'no'
                            fail_count += 1
                            try: git.execute(['git', 'cherry-pick', '--abort'])
                            except GitCommandError: pass
                    else:
                        print("⏭️  Skipping commit.")
                        results[commit_id] = 'no'
                        fail_count += 1
                        try: git.execute(['git', 'cherry-pick', '--abort'])
                        except GitCommandError: pass
                else:
                    print(f"❌ Conflict. Aborting this commit.")
                    results[commit_id] = 'no'
                    fail_count += 1
                    try: git.execute(['git', 'cherry-pick', '--abort'])
                    except GitCommandError: pass

    # === SAVE RESULTS ===
    print(f"\nProcess completed. Success: {success_count}, Failed: {fail_count}, Already Present: {empty_count}")
    print(f"Updating {excel_file} while preserving hyperlinks...")
    
    try:
        wb = load_workbook(excel_file)
        ws = wb.active
        col_map = {cell.value: cell.column for cell in ws[1]}
        
        # Ensure 'Result' column exists
        if 'Result' not in col_map:
            new_col = ws.max_column + 1
            ws.cell(row=1, column=new_col).value = 'Result'
            col_map['Result'] = new_col
            
        success_col = col_map['Result']
        id_col = col_map['Commit SHA']
        in_release_col = col_map.get('Audit Trail')
        decision_col = col_map.get('Action')

        for row_idx in range(2, ws.max_row + 1):
            commit_id = ws.cell(row=row_idx, column=id_col).value
            if commit_id in results:
                status = results[commit_id]
                ws.cell(row=row_idx, column=success_col).value = status
                if status == 'no changes to commit':
                    if in_release_col: ws.cell(row=row_idx, column=in_release_col).value = 'Yes'
                    if decision_col: ws.cell(row=row_idx, column=decision_col).value = 'no'

        wb.save(excel_file)
        print(f"Successfully updated {excel_file}.")
    except Exception as e:
        print(f"❌ Error updating Excel: {e}")

if __name__ == "__main__":
    main()
