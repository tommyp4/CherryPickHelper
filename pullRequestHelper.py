import os, re
import pandas as pd
from datetime import datetime, timedelta
from azure.devops.connection import Connection
from msrest.authentication import BasicAuthentication
from azure.devops.v7_1.git.models import GitQueryCommitsCriteria, GitVersionDescriptor, GitPullRequestSearchCriteria, GitPullRequestQuery, GitPullRequestQueryInput
import config
from shared import get_matched_config_author

def get_first_line(msg):
    if not msg: return ""
    return msg.split('\n')[0].strip()

def clean_subject(subject):
    # 1. Strip system merge noise
    subject = re.sub(r'^Merged PR \d+: ', '', subject, flags=re.IGNORECASE)
    # 2. Strip "Cherry-pick" noise from UI-based cherry picks
    subject = re.sub(r'^Cherry-?pick\s+"?', '', subject, flags=re.IGNORECASE)
    subject = re.sub(r'"?\s+into\s+release/.*$', '', subject, flags=re.IGNORECASE)
    subject = re.sub(r'"?\s+into\s+develop.*$', '', subject, flags=re.IGNORECASE)
    # 3. Strip Jira IDs
    subject = re.sub(r'ALP[-_]?\d+', '', subject, flags=re.IGNORECASE)
    # 4. Strip "Revert" and cleanup
    subject = re.sub(r'^Revert\s+"?', '', subject, flags=re.IGNORECASE)
    subject = re.sub(r'[^\w\s]', '', subject)
    return subject.strip().lower()

def extract_jira_id(text):
    match = re.search(r'ALP[-_]?(\d+)', str(text), re.IGNORECASE)
    if match:
        return f"ALP-{match.group(1)}"
    return None

def extract_original_pr_id(text):
    """Looks for 'cherry picked from !1234' in commit message."""
    if not text: return None
    match = re.search(r'cherry[- ]picked from !(\d+)', text, re.IGNORECASE)
    if match:
        res = match.group(1)
        # print(f"DEBUG: Found PR link !{res} in text: {text[:50]}...")
        return res
    return None

def extract_original_commit_id(text):
    """Looks for 'cherry-picked from commit a8feecf1' in commit message."""
    if not text: return None
    match = re.search(r'cherry[- ]picked from commit `?([a-f0-9]{7,40})`?', text, re.IGNORECASE)
    if match:
        res = match.group(1)
        # print(f"DEBUG: Found Commit link {res} in text: {text[:50]}...")
        return res
    return None

def get_pr_ids_for_commits(git_client, repo_id, project, commit_ids):
    """Batch queries Azure DevOps for PRs associated with commit IDs."""
    results = {}
    commit_ids = list(set(commit_ids))
    total_found = 0
    # Process in chunks of 10 (Azure DevOps API limit)
    for i in range(0, len(commit_ids), 10):
        chunk = [cid.lower() for cid in commit_ids[i:i+10]]
        try:
            query = GitPullRequestQuery(
                queries=[GitPullRequestQueryInput(items=[cid], type='commit') for cid in chunk]
            )
            resp = git_client.get_pull_request_query(query, repo_id, project=project)
            if resp and resp.results:
                for mapping in resp.results:
                    for cid, prs in mapping.items():
                        if prs:
                            # Normalize key to lowercase
                            results[cid.lower()] = {str(pr.pull_request_id) for pr in prs}
                            total_found += 1
        except Exception as e:
            pass # Silent fail for individual chunks
    print(f"  PR Context: Found PRs for {total_found} out of {len(commit_ids)} commits.")
    return results

def get_commit_files(git_client, repo_id, commit_id, project):
    """Fetches the list of file paths changed in a commit."""
    try:
        changes = git_client.get_changes(commit_id, repo_id, project=project)
        return {c.item.path for c in changes.changes if c.item and c.item.path}
    except Exception:
        return set()

def extract_all_pr_ids(text):
    """Extracts all PR IDs (e.g. !1234 or PR 1234) from text."""
    if not text: return set()
    # Find !1234
    ids = set(re.findall(r'!(\d+)', text))
    # Find PR 1234
    ids.update(re.findall(r'PR\s?(\d+)', text, re.IGNORECASE))
    return ids

def main():
    # === AUTHENTICATION ===
    credentials = BasicAuthentication('', config.personal_access_token)
    connection = Connection(base_url=config.organization_url, creds=credentials)
    git_client = connection.clients.get_git_client()
    repo = git_client.get_repository(project=config.project_name, repository_id=config.repository_name)
    
    start_dt = datetime.strptime(config.start_date, '%Y-%m-%d %H:%M')

    # === 1. SCAN RELEASE HISTORIES (Building Inventories) ===
    release_global_inventory = {} # (author, cleaned) -> {'count': 0, 'prs': set()}
    release_jira_inventory = {}   # jira_id -> {cleaned: {'count': 0, 'prs': set()}}
    release_linked_prs = {} # { original_pr_id: release_pr_id }
    release_linked_commits = set()
    all_release_cleaned_subjects = set()

    branches_to_check = [config.check_target_branch, config.cherry_pick_base_branch]
    pr_description_cache = {} # { pr_id: description_text }

    for branch_name in branches_to_check:
        print(f"Inventorying '{branch_name}' starting from {config.start_date}...")
        try:
            ver = GitVersionDescriptor(version=branch_name, version_type='branch')
            crit = GitQueryCommitsCriteria(item_version=ver, from_date=config.start_date)
            
            b_global = {} # (author, cleaned) -> {'count': 0, 'prs': set()}
            b_jira = {}   # jira_id -> {cleaned: {'count': 0, 'prs': set()}}

            commits = git_client.get_commits(repo.id, search_criteria=crit, project=config.project_name, top=1000)
            if commits:
                print(f"  Processing {len(commits)} commits in {branch_name} (fetching full details)...")
                for c in commits:
                    try:
                        # Always fetch full details to ensure we get the complete message body/comment
                        full_c = git_client.get_commit(c.commit_id, repo.id, project=config.project_name)
                        msg = full_c.comment
                    except Exception:
                        msg = c.comment # Fallback if full fetch fails

                    subj = get_first_line(msg)
                    cleaned = clean_subject(subj)
                    author = get_matched_config_author(c.author.name, config.authors)
                    jira = extract_jira_id(msg)

                    # Extract Release PR ID from subject (e.g., "Merged PR 1234: ...")
                    rel_pr_match = re.search(r'^Merged PR (\d+):', subj, re.IGNORECASE)
                    release_pr_id = rel_pr_match.group(1) if rel_pr_match else None

                    # DEEP PR SCANNING: Fetch description from API if we found a PR ID
                    pr_body_links = set()
                    if release_pr_id:
                        if release_pr_id not in pr_description_cache:
                            try:
                                pr_detail = git_client.get_pull_request(repo.id, int(release_pr_id), project=config.project_name)
                                pr_description_cache[release_pr_id] = pr_detail.description if pr_detail.description else ""
                            except Exception:
                                pr_description_cache[release_pr_id] = ""
                        
                        pr_body_links = extract_all_pr_ids(pr_description_cache[release_pr_id])

                    # High-fidelity link extraction (Combine commit msg + PR body links)
                    found_pr_ids = extract_all_pr_ids(msg)
                    found_pr_ids.update(pr_body_links)

                    for pid in found_pr_ids:
                        if pid != release_pr_id: # Don't map it back to itself
                            release_linked_prs[pid] = release_pr_id
                    
                    linked_commit = extract_original_commit_id(msg)
                    if linked_commit: release_linked_commits.add(linked_commit)

                    if cleaned:
                        all_release_cleaned_subjects.add(cleaned)
                        if author:
                            key = (author, cleaned)
                            if key not in b_global: b_global[key] = {'count': 0, 'prs': set()}
                            b_global[key]['count'] += 1
                            if release_pr_id: b_global[key]['prs'].add(release_pr_id)
                        if jira:
                            if jira not in b_jira: b_jira[jira] = {}
                            if cleaned not in b_jira[jira]: b_jira[jira][cleaned] = {'count': 0, 'prs': set()}
                            b_jira[jira][cleaned]['count'] += 1
                            if release_pr_id: b_jira[jira][cleaned]['prs'].add(release_pr_id)
            
            # Merge branch results into global inventory
            for k, v in b_global.items():
                if k not in release_global_inventory: release_global_inventory[k] = {'count': 0, 'prs': set()}
                release_global_inventory[k]['count'] = max(release_global_inventory[k]['count'], v['count'])
                release_global_inventory[k]['prs'].update(v['prs'])
            for jid, sub_counts in b_jira.items():
                if jid not in release_jira_inventory: release_jira_inventory[jid] = {}
                for s, v in sub_counts.items():
                    if s not in release_jira_inventory[jid]: release_jira_inventory[jid][s] = {'count': 0, 'prs': set()}
                    release_jira_inventory[jid][s]['count'] = max(release_jira_inventory[jid][s]['count'], v['count'])
                    release_jira_inventory[jid][s]['prs'].update(v['prs'])
        except Exception as e:
            print(f"Warning: Could not scan {branch_name}: {e}")

    print(f"Found {len(release_linked_prs)} unique PR links and {len(release_linked_commits)} commit links in release history.")
    # === 2. SCAN DEVELOP (Phase A: Collection & Dependency Analysis) ===
    print(f"Scanning history of '{config.main_target_branch}'...")
    dev_ver = GitVersionDescriptor(version=config.main_target_branch, version_type='branch')
    dev_crit = GitQueryCommitsCriteria(item_version=dev_ver)

    team_commits = []
    foreign_commits = []
    team_files = set() # Set of all files touched by the team

    skip = 0
    top = 100
    while skip < 1500: # Slightly smaller window for performance
        history_commits = git_client.get_commits(repo.id, search_criteria=dev_crit, project=config.project_name, top=top, skip=skip)
        if not history_commits: break
            
        for commit in history_commits:
            commit_date = commit.committer.date.replace(tzinfo=None)
            if commit_date < start_dt:
                if skip > 300: skip = 1500; break
                continue

            try:
                # Always fetch full details to ensure we get the complete message body/comment
                full_commit = git_client.get_commit(commit.commit_id, repo.id, project=config.project_name)
                msg = full_commit.comment.strip()
            except Exception:
                msg = commit.comment.strip() # Fallback

            subject = get_first_line(msg)
            
            # Skip structural PR merge records
            if subject.startswith("Merged PR"):
                if full_commit.parents and len(full_commit.parents) > 1: continue

            matched_author = get_matched_config_author(commit.author.name, config.authors)
            
            if matched_author:
                # print(f"  Team commit: {commit.commit_id[:8]} by {matched_author}")
                # Fetch files to build dependency map
                files = get_commit_files(git_client, repo.id, commit.commit_id, config.project_name)
                team_files.update(files)
                
                team_commits.append({
                    'commit': commit,
                    'author': matched_author,
                    'subject': subject,
                    'cleaned': clean_subject(subject),
                    'jira': extract_jira_id(msg),
                    'date': commit_date,
                    'full_msg': msg,
                    'files': files,
                    'type': 'Team'
                })
            else:
                foreign_commits.append(commit)

        skip += top

    # === 3. FILTER FOREIGN COMMITS FOR OVERLAPS ===
    print(f"Checking {len(foreign_commits)} foreign commits for file overlaps...")
    overlapping_foreign = []
    for c in foreign_commits:
        f_files = get_commit_files(git_client, repo.id, c.commit_id, config.project_name)
        overlaps = f_files.intersection(team_files)
        if overlaps:
            print(f"  🔗 Overlap found: {c.commit_id[:8]} by {c.author.name} touches team files!")
            msg = c.comment.strip()
            overlapping_foreign.append({
                'commit': c,
                'author': f"FOREIGN: {c.author.name}",
                'subject': get_first_line(msg),
                'cleaned': clean_subject(get_first_line(msg)),
                'jira': extract_jira_id(msg),
                'date': c.committer.date.replace(tzinfo=None),
                'full_msg': msg,
                'type': 'Dependency',
                'overlap_files': list(overlaps)[:3] # Show first few overlaps
            })

    # Combine and sort
    all_develop_work = team_commits + overlapping_foreign
    all_develop_work.sort(key=lambda x: x['date'])

    # Recalculate develop counts for combined list
    dev_global_counts = {}
    dev_jira_counts = {}
    all_dev_commit_ids = [item['commit'].commit_id for item in all_develop_work]
    
    print(f"Fetching PR context for {len(all_dev_commit_ids)} develop commits...")
    commit_to_pr_map = get_pr_ids_for_commits(git_client, repo.id, config.project_name, all_dev_commit_ids)

    for item in all_develop_work:
        if item['type'] == 'Team':
            g_key = (item['author'], item['cleaned'])
            dev_global_counts[g_key] = dev_global_counts.get(g_key, 0) + 1
            if item['jira']:
                jid = item['jira']
                if jid not in dev_jira_counts: dev_jira_counts[jid] = {}
                dev_jira_counts[jid][item['cleaned']] = dev_jira_counts[jid].get(item['cleaned'], 0) + 1

    # === 4. DECISION LOGIC ===
    print(f"Analyzing {len(all_develop_work)} total relevant commits...")
    all_final_rows = []
    
    match_counts = {"PR Link": 0, "Commit Link": 0, "Exact": 0, "Ticket": 0, "No": 0}

    for item in all_develop_work:
        is_in_release = "No"
        author = item['author']
        cleaned = item['cleaned']
        jira = item['jira']
        commit_id = item['commit'].commit_id
        
        if item['type'] == 'Dependency':
            is_in_release = f"Foreign commit touching {item['overlap_files'][0]}"
            decision = "" # Leave blank for review
        else:
            # 1. High-fidelity Link Matching (Highest Confidence)
            # Check for direct commit ID match (short or long)
            matched_linked_commit = next((c for c in release_linked_commits if commit_id.lower().startswith(c.lower()) or c.lower().startswith(commit_id.lower())), None)
            
            # Check for PR ID match
            my_prs = commit_to_pr_map.get(commit_id.lower(), set())
            matched_linked_pr = next((p for p in my_prs if p in release_linked_prs), None)

            if matched_linked_commit:
                is_in_release = f"Yes (Commit Link Match: {matched_linked_commit[:8]})"
                match_counts["Commit Link"] += 1
            elif matched_linked_pr:
                is_in_release = "Yes (PR Link Match)"
                match_counts["PR Link"] += 1
            else:
                # 2. Standard Heuristic Rules
                g_key = (author, cleaned)
                if g_key in release_global_inventory:
                    if dev_global_counts[g_key] == release_global_inventory[g_key]['count']:
                        is_in_release = "Yes (Exact Match)"
                        match_counts["Exact"] += 1
                    else: is_in_release = "Needs attention (Subject Count Mismatch)"
                elif jira and jira in release_jira_inventory:
                    if cleaned in release_jira_inventory[jira]:
                        if dev_jira_counts[jira][cleaned] == release_jira_inventory[jira][cleaned]['count']:
                            is_in_release = "Yes (Ticket Match)"
                            match_counts["Ticket"] += 1
                        else: is_in_release = "Needs attention (Ticket Count Mismatch)"
                    else: is_in_release = f"Likely (Jira {jira} exists, but subjects differ)"
                
                if is_in_release == "No":
                    match_counts["No"] += 1
            
            decision = "no" if is_in_release.startswith("Yes") else "yes" if is_in_release == "No" else ""

        jira_url = f"{config.jira_base_url}{jira}" if jira else None
        
        # Determine PR ID to link (Always link if known, even if not matched in release)
        my_prs = commit_to_pr_map.get(commit_id.lower(), set())
        # Prioritize the matched release PR if it exists, otherwise just the first PR found
        display_pr = next((p for p in my_prs if p in release_linked_prs), (sorted(list(my_prs))[0] if my_prs else None))
        pr_url = f"{config.organization_url}/{config.project_name}/_git/{config.repository_name}/pullrequest/{display_pr}" if display_pr else None

        # Link to the Matched Release PR
        # 1. Try explicit link first
        matched_release_pr_id = release_linked_prs.get(display_pr) if display_pr in release_linked_prs else None
        
        # 2. Fallback to inferred links from inventories if matched
        if not matched_release_pr_id:
            g_key = (author, cleaned)
            if is_in_release == "Yes (Exact Match)" and g_key in release_global_inventory:
                prs = release_global_inventory[g_key]['prs']
                if prs: matched_release_pr_id = sorted(list(prs))[0]
            elif is_in_release == "Yes (Ticket Match)" and jira in release_jira_inventory:
                if cleaned in release_jira_inventory[jira]:
                    prs = release_jira_inventory[jira][cleaned]['prs']
                    if prs: matched_release_pr_id = sorted(list(prs))[0]
        
        matched_rel_url = f"{config.organization_url}/{config.project_name}/_git/{config.repository_name}/pullrequest/{matched_release_pr_id}" if matched_release_pr_id else None

        all_final_rows.append({
            'Jira Link': jira,
            'Jira URL': jira_url,
            'Description': item['subject'],
            'Audit Trail': is_in_release,
            'Jira Target': "", # Placeholder for jiraHelper
            'Action': decision,
            'Result': "", # Placeholder for cherryPickHelper
            'Original PR': f"!{display_pr}" if display_pr else "",
            'PR URL': pr_url,
            'Release PR': f"!{matched_release_pr_id}" if matched_release_pr_id else "",
            'Matched Rel URL': matched_rel_url,
            'Owner': author,
            'Merged Date': item['date'],
            'Commit SHA': item['commit'].commit_id,
            'Full Message': item['full_msg']
        })

    print(f"\nDecision Summary:")
    for k, v in match_counts.items():
        print(f"  - {k}: {v}")

    # === SAVE TO EXCEL (Using openpyxl for pipeline compatibility) ===
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Commits"

    # Write Headers (Human-First Order)
    headers = [
        'Jira ID', 'Description', 'Audit Trail', 'Code Present?', 
        'Jira Target', 'Action', 'Result', 'Original PR', 
        'Release PR', 'Owner', 'Merged Date', 'Commit SHA', 'Full Message'
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col).value = header

    # Write Rows
    for row_idx, data in enumerate(all_final_rows, 2):
        # 1. Jira ID (Hyperlinked)
        c1 = ws.cell(row=row_idx, column=1)
        c1.value = data['Jira Link']
        if data['Jira URL']:
            c1.hyperlink = data['Jira URL']
            c1.font = Font(color="0000FF", underline="single")

        # 2. Description
        ws.cell(row=row_idx, column=2).value = data['Description']

        # 3. Audit Trail (PR Links, etc.)
        ws.cell(row=row_idx, column=3).value = data['Audit Trail']

        # 4. Code Present? (Placeholder for verifyBranchSync)
        ws.cell(row=row_idx, column=4).value = ""

        # 5. Jira Target (Placeholder for jiraHelper)
        ws.cell(row=row_idx, column=5).value = data['Jira Target']

        # 6. Action (Decision)
        ws.cell(row=row_idx, column=6).value = data['Action']

        # 7. Result (Success/Fail)
        ws.cell(row=row_idx, column=7).value = data['Result']

        # 8. Original PR (Hyperlinked)
        c8 = ws.cell(row=row_idx, column=8)
        c8.value = data['Original PR']
        if data['PR URL']:
            c8.hyperlink = data['PR URL']
            c8.font = Font(color="0000FF", underline="single")

        # 9. Release PR (Hyperlinked)
        c9 = ws.cell(row=row_idx, column=9)
        c9.value = data['Release PR']
        if data['Matched Rel URL']:
            c9.hyperlink = data['Matched Rel URL']
            c9.font = Font(color="0000FF", underline="single")

        # 10. Owner
        ws.cell(row=row_idx, column=10).value = data['Owner']

        # 11. Merged Date
        ws.cell(row=row_idx, column=11).value = data['Merged Date']

        # 12. Commit SHA
        ws.cell(row=row_idx, column=12).value = data['Commit SHA']

        # 13. Full Message
        ws.cell(row=row_idx, column=13).value = data['Full Message']

    excel_file = 'cherrypick_list.xlsx'
    try:
        wb.save(excel_file)
        print(f"Successfully saved {len(all_final_rows)} commits to {excel_file}")
    except Exception:
        print(f"❌ Error: Close the file {excel_file}.")

if __name__ == "__main__":
    main()
